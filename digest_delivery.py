from __future__ import annotations

import json
import os
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

from docx import Document as create_document
from docx.document import Document as DocxDocument


@dataclass
class DigestDocumentResult:
    docx_path: str | None
    total_items: int
    group_count: int
    group_sizes: dict[str, int]


class DailyDigestDocumentService:
    def __init__(self, *, db, reports_dir: str):
        self.db = db
        self.reports_dir = reports_dir

    async def build_document(
        self,
        *,
        hours: int = 24,
        group_by: str = "niche",
        min_wtp: int = 8,
        max_items_per_group: int = 10,
    ) -> DigestDocumentResult:
        rows = await self.db.get_recent_pain_points(hours=hours, limit=500)
        filtered_rows = [row for row in rows if self._include_row(row, min_wtp=min_wtp)]
        if not filtered_rows:
            return DigestDocumentResult(docx_path=None, total_items=0, group_count=0, group_sizes={})

        filtered_post_ids = [str(row.get("post_id")) for row in filtered_rows if row.get("post_id")]
        canonical_clusters = await self.db.get_latest_canonical_clusters(limit=6, post_ids=filtered_post_ids)

        current_rows = [row for row in filtered_rows if self._opportunity_bucket(row) == "current_opportunity"]
        evergreen_rows = [row for row in filtered_rows if self._opportunity_bucket(row) == "evergreen_pain"]
        unknown_rows = [row for row in filtered_rows if self._opportunity_bucket(row) == "unknown_age"]
        current_groups = self._order_groups(self._group_rows(current_rows, group_by=group_by)) if current_rows else []
        evergreen_groups = self._order_groups(self._group_rows(evergreen_rows, group_by=group_by)) if evergreen_rows else []
        unknown_groups = self._order_groups(self._group_rows(unknown_rows, group_by=group_by)) if unknown_rows else []

        document = create_document()
        document.add_heading("Pain Finder Daily Digest", level=0)
        document.add_paragraph(
            (
                f"Generated: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}\n"
                f"Window: last {hours}h\n"
                f"Grouping: {group_by}\n"
                f"Current opportunities: {sum(len(items) for _, items in current_groups)}\n"
                f"Evergreen pain index: {sum(len(items) for _, items in evergreen_groups)}\n"
                f"Unknown age review queue: {sum(len(items) for _, items in unknown_groups)}"
            )
        )

        overview = document.add_paragraph()
        overview.add_run("Current opportunity groups: ").bold = True
        overview.add_run(", ".join(f"{label} ({len(items)})" for label, items in current_groups[:8]) or "none")

        evergreen_overview = document.add_paragraph()
        evergreen_overview.add_run("Evergreen groups: ").bold = True
        evergreen_overview.add_run(", ".join(f"{label} ({len(items)})" for label, items in evergreen_groups[:8]) or "none")

        unknown_overview = document.add_paragraph()
        unknown_overview.add_run("Unknown-age groups: ").bold = True
        unknown_overview.add_run(", ".join(f"{label} ({len(items)})" for label, items in unknown_groups[:8]) or "none")

        blockers = self._recurring_blockers(filtered_rows)
        if blockers:
            blockers_paragraph = document.add_paragraph()
            blockers_paragraph.add_run("Recurring blockers: ").bold = True
            blockers_paragraph.add_run(" | ".join(blockers))

        if canonical_clusters:
            document.add_heading("Canonical pain clusters", level=1)
            self._render_cluster_section(document, canonical_clusters)

        if current_groups:
            document.add_heading("Current opportunities", level=1)
            self._render_grouped_section(document, current_groups, max_items_per_group=max_items_per_group)
        if evergreen_groups:
            document.add_heading("Evergreen pain index", level=1)
            self._render_grouped_section(document, evergreen_groups, max_items_per_group=max_items_per_group)
        if unknown_groups:
            document.add_heading("Unknown age review queue", level=1)
            self._render_grouped_section(document, unknown_groups, max_items_per_group=max_items_per_group)

        os.makedirs(self.reports_dir, exist_ok=True)
        timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S_%f")
        final_path = os.path.join(self.reports_dir, f"daily_digest_{timestamp}.docx")
        tmp_path = f"{final_path}.tmp"
        try:
            document.save(tmp_path)
            os.replace(tmp_path, final_path)
        except Exception:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
            raise

        all_groups = current_groups + evergreen_groups + unknown_groups
        return DigestDocumentResult(
            docx_path=final_path,
            total_items=sum(len(items[:max_items_per_group]) for _, items in all_groups),
            group_count=len(all_groups),
            group_sizes={label: len(items) for label, items in all_groups},
        )

    @staticmethod
    def _include_row(row: dict[str, Any], *, min_wtp: int) -> bool:
        triage_status = str(row.get("triage_status") or "new").strip().lower()
        if triage_status == "favorite":
            return True
        return int(row.get("willingness_to_pay") or 0) >= min_wtp

    @staticmethod
    def _group_label(row: dict[str, Any], *, group_by: str) -> str:
        if group_by == "source":
            return (str(row.get("source") or "unknown").strip() or "unknown").title()
        if group_by == "category":
            return (str(row.get("category") or "Uncategorized").strip() or "Uncategorized").title()
        return str(row.get("niche_category") or "Uncategorized").strip() or "Uncategorized"

    def _group_rows(self, rows: list[dict[str, Any]], *, group_by: str) -> dict[str, list[dict[str, Any]]]:
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            grouped[self._group_label(row, group_by=group_by)].append(row)
        return grouped

    @staticmethod
    def _opportunity_bucket(row: dict[str, Any]) -> str:
        bucket = str(row.get("opportunity_bucket") or "").strip().lower()
        if bucket == "current_opportunity":
            return "current_opportunity"
        if bucket == "evergreen_pain":
            return "evergreen_pain"
        return "unknown_age"

    def _render_cluster_section(self, document: DocxDocument, clusters: list[dict[str, Any]]) -> None:
        for cluster in clusters:
            label = (str(cluster.get("label") or "Recurring pain cluster").strip() or "Recurring pain cluster")
            summary = (str(cluster.get("summary") or "No summary available.").strip() or "No summary available.")
            avg_score = float(cluster.get("avg_opportunity_score") or 0.0)
            fresh_post_count = int(cluster.get("fresh_post_count") or 0)
            evergreen_post_count = int(cluster.get("evergreen_post_count") or 0)
            incumbents = cluster.get("incumbents") or []
            incumbents_text = ", ".join(str(item) for item in incumbents[:4]) or "none"

            document.add_heading(label, level=2)
            metrics = document.add_paragraph(
                f"Avg opp {avg_score:.1f} | Fresh {fresh_post_count} | Evergreen {evergreen_post_count}"
            )
            metrics.style = "Intense Quote"
            document.add_paragraph(summary)
            document.add_paragraph(f"Dominant incumbents: {incumbents_text}")

    def _render_grouped_section(
        self,
        document: DocxDocument,
        ordered_groups: list[tuple[str, list[dict[str, Any]]]],
        *,
        max_items_per_group: int,
    ) -> None:
        for label, items in ordered_groups:
            document.add_heading(f"{label} ({len(items)})", level=2)
            for row in items[:max_items_per_group]:
                title = (row.get("title") or "Untitled").strip() or "Untitled"
                summary = (row.get("summary") or "No summary available.").strip() or "No summary available."
                source = (row.get("source") or "unknown").strip() or "unknown"
                subreddit = (row.get("subreddit") or "n/a").strip() or "n/a"
                url = (row.get("url") or "").strip()
                wtp = int(row.get("willingness_to_pay") or 0)
                pain_level = int(row.get("pain_level") or 0)
                opportunity_score = self._row_opportunity_score(row)
                competitors = ", ".join(self._competitor_tags(row)) or "none"
                deep_dive_summary = (row.get("deep_dive_summary") or "").strip()
                promotion_rejection_reason = self._promotion_rejection_reason(row)

                header = document.add_paragraph()
                header.add_run(title).bold = True
                metrics = document.add_paragraph(
                    f"Opp {opportunity_score:.1f} | WTP {wtp}/10 | Pain {pain_level}/10 | Source {source} | Scope {subreddit}"
                )
                metrics.style = "Intense Quote"
                document.add_paragraph(summary)
                document.add_paragraph(f"Competitors/tags: {competitors}")
                if promotion_rejection_reason:
                    document.add_paragraph(f"Promotion blocked: {promotion_rejection_reason}")
                if deep_dive_summary:
                    document.add_paragraph(f"Deep dive: {deep_dive_summary}")
                if url:
                    document.add_paragraph(f"Link: {url}")

    @classmethod
    def _order_groups(cls, grouped: dict[str, list[dict[str, Any]]]) -> list[tuple[str, list[dict[str, Any]]]]:
        def _group_score(item: tuple[str, list[dict[str, Any]]]) -> tuple[float, int, str]:
            label, rows = item
            best_score = max(cls._row_opportunity_score(row) for row in rows)
            return (-best_score, -len(rows), label.lower())

        ordered_groups = sorted(grouped.items(), key=_group_score)
        for _, rows in ordered_groups:
            rows.sort(
                key=lambda row: (
                    cls._row_opportunity_score(row),
                    int(row.get("source_created_ts") or 0),
                    int(row.get("willingness_to_pay") or 0),
                    int(row.get("pain_level") or 0),
                ),
                reverse=True,
            )
        return ordered_groups

    @staticmethod
    def _row_opportunity_score(row: dict[str, Any]) -> float:
        raw = row.get("opportunity_score")
        try:
            if raw is not None:
                return float(raw)
        except (TypeError, ValueError):
            pass
        return float(int(row.get("willingness_to_pay") or 0))

    @staticmethod
    def _competitor_tags(row: dict[str, Any]) -> list[str]:
        raw = row.get("competitor_tags") or []
        if isinstance(raw, list):
            return [str(item) for item in raw if str(item).strip()]
        if isinstance(raw, str):
            try:
                parsed = json.loads(raw)
            except json.JSONDecodeError:
                parsed = [raw]
            if isinstance(parsed, list):
                return [str(item) for item in parsed if str(item).strip()]
            if parsed:
                return [str(parsed)]
        return []

    @staticmethod
    def _analysis_payload(row: dict[str, Any]) -> dict[str, Any]:
        raw = row.get("analysis_payload_json") or row.get("analysis_payload") or {}
        if isinstance(raw, dict):
            return raw
        if isinstance(raw, str) and raw.strip():
            try:
                parsed = json.loads(raw)
            except json.JSONDecodeError:
                return {}
            return parsed if isinstance(parsed, dict) else {}
        return {}

    @classmethod
    def _promotion_rejection_reason(cls, row: dict[str, Any]) -> str:
        raw_reason = row.get("evidence_rejection_reason")
        if raw_reason is None:
            raw_reason = cls._analysis_payload(row).get("evidence_rejection_reason")
        return str(raw_reason or "").strip()

    @staticmethod
    def _recurring_blockers(rows: list[dict[str, Any]]) -> list[str]:
        counts: dict[str, int] = defaultdict(int)
        for row in rows:
            summary = str(row.get("deep_dive_summary") or "").strip()
            if summary:
                counts[summary] += 1
        return [summary for summary, _ in sorted(counts.items(), key=lambda item: item[1], reverse=True)[:5]]
